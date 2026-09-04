from dataclasses import dataclass
from typing import List
from pathlib import Path
import pypdf
import docx
from bs4 import BeautifulSoup
import csv
import markdown

@dataclass
class ParsedPage:
    page_num: int
    text: str
    metadata: dict

@dataclass
class ParseResult:
    full_text: str
    pages: List[ParsedPage]
    page_count: int
    language: str
    metadata: dict

async def parse_document(file_path: Path, mime_type: str) -> ParseResult:
    if mime_type == "application/pdf":
        return await _parse_pdf(file_path)
    elif mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return await _parse_docx(file_path)
    elif mime_type in ("text/plain", "text/markdown"):
        return await _parse_text(file_path)
    elif mime_type == "text/csv":
        return await _parse_csv(file_path)
    elif mime_type in ("text/html", "application/xhtml+xml"):
        return await _parse_html(file_path)
    else:
        raise ValueError(f"Unsupported MIME type: {mime_type}")

async def _parse_pdf(file_path: Path) -> ParseResult:
    pages = []
    full_text_parts = []
    
    with pypdf.PdfReader(file_path) as reader:
        for i, page in enumerate(reader.pages):
            text = page.extract_text()
            if text:
                pages.append(ParsedPage(
                    page_num=i + 1,
                    text=text,
                    metadata={"source": "pdf", "page": i + 1},
                ))
                full_text_parts.append(f"[Page {i+1}]\n{text}")
    
    if not full_text_parts:
        return ParseResult(
            full_text="",
            pages=[],
            page_count=len(reader.pages),
            language="unknown",
            metadata={"scanned": True, "warning": "No extractable text - OCR required"},
        )
    
    return ParseResult(
        full_text="\n\n".join(full_text_parts),
        pages=pages,
        page_count=len(reader.pages),
        language=_detect_language(full_text_parts[0][:1000]),
        metadata={"parser": "pypdf"},
    )

async def _parse_docx(file_path: Path) -> ParseResult:
    doc = docx.Document(file_path)
    pages = []
    full_text_parts = []
    
    text_parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            text_parts.append(para.text)
    
    for table in doc.tables:
        table_text = []
        for row in table.rows:
            row_text = [cell.text for cell in row.cells]
            table_text.append(" | ".join(row_text))
        if table_text:
            text_parts.append("\n".join(table_text))
    
    full_text = "\n\n".join(text_parts)
    pages.append(ParsedPage(
        page_num=1,
        text=full_text,
        metadata={"source": "docx", "paragraphs": len(doc.paragraphs), "tables": len(doc.tables)},
    ))
    
    return ParseResult(
        full_text=full_text,
        pages=pages,
        page_count=1,
        language=_detect_language(full_text[:1000]),
        metadata={"parser": "python-docx"},
    )

async def _parse_text(file_path: Path) -> ParseResult:
    text = file_path.read_text(encoding="utf-8", errors="replace")
    return ParseResult(
        full_text=text,
        pages=[ParsedPage(page_num=1, text=text, metadata={"source": "text"})],
        page_count=1,
        language=_detect_language(text[:1000]),
        metadata={"parser": "text"},
    )

async def _parse_csv(file_path: Path) -> ParseResult:
    rows = []
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        reader = csv.reader(f)
        headers = next(reader, [])
        for i, row in enumerate(reader):
            if any(cell.strip() for cell in row):
                row_dict = dict(zip(headers, row))
                rows.append(" | ".join(f"{k}: {v}" for k, v in row_dict.items()))
    
    full_text = "\n".join(rows)
    return ParseResult(
        full_text=full_text,
        pages=[ParsedPage(page_num=1, text=full_text, metadata={"source": "csv", "rows": len(rows), "columns": headers})],
        page_count=1,
        language=_detect_language(full_text[:1000]),
        metadata={"parser": "csv", "columns": headers},
    )

async def _parse_html(file_path: Path) -> ParseResult:
    html = file_path.read_text(encoding="utf-8", errors="replace")
    soup = BeautifulSoup(html, "html.parser")
    
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()
    
    text = soup.get_text(separator="\n", strip=True)
    
    return ParseResult(
        full_text=text,
        pages=[ParsedPage(page_num=1, text=text, metadata={"source": "html", "title": soup.title.string if soup.title else ""})],
        page_count=1,
        language=_detect_language(text[:1000]),
        metadata={"parser": "beautifulsoup4"},
    )

def _detect_language(text: str) -> str:
    if not text:
        return "unknown"
    english_words = {"the", "and", "is", "to", "of", "a", "in", "that", "it", "for"}
    words = set(text.lower().split()[:100])
    if len(words & english_words) > 3:
        return "en"
    return "unknown"